import os
import re
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def create_docx_from_markdown(md_path, docx_path):
    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    doc = docx.Document()

    # Page setup - Margins
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Styles setup
    style_normal = doc.styles['Normal']
    font = style_normal.font
    font.name = 'Calibri'
    font.size = Pt(11)
    font.color.rgb = RGBColor(0x22, 0x22, 0x22)

    in_code_block = False
    code_lines = []
    in_table = False
    table_rows = []

    i = 0
    n = len(lines)

    while i < n:
        line = lines[i].rstrip('\r\n')

        # Code block handling
        if line.startswith('```'):
            if in_code_block:
                # Flush code block
                in_code_block = False
                code_text = '\n'.join(code_lines)
                table = doc.add_table(rows=1, cols=1)
                table.alignment = WD_TABLE_ALIGNMENT.CENTER
                cell = table.cell(0, 0)
                set_cell_background(cell, "F4F6F8")
                set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
                p = cell.paragraphs[0]
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.05
                run = p.add_run(code_text)
                run.font.name = 'Consolas'
                run.font.size = Pt(9)
                run.font.color.rgb = RGBColor(0x1F, 0x24, 0x28)
                doc.add_paragraph() # Spacing
                code_lines = []
            else:
                in_code_block = True
                code_lines = []
            i += 1
            continue

        if in_code_block:
            code_lines.append(line)
            i += 1
            continue

        # Table handling
        if '|' in line and (line.strip().startswith('|') or line.strip().endswith('|')):
            table_rows.append(line)
            # Peek if table continues
            if i + 1 < n and '|' in lines[i + 1]:
                i += 1
                continue
            else:
                # Process table
                parsed_rows = []
                for tline in table_rows:
                    if re.match(r'^\s*\|?\s*[-:]+[-| :]*\|?\s*$', tline):
                        continue # separator line
                    cells = [c.strip() for c in tline.strip().strip('|').split('|')]
                    parsed_rows.append(cells)

                if parsed_rows:
                    col_count = max(len(r) for r in parsed_rows)
                    t = doc.add_table(rows=len(parsed_rows), cols=col_count)
                    t.alignment = WD_TABLE_ALIGNMENT.CENTER
                    for r_idx, row_data in enumerate(parsed_rows):
                        for c_idx in range(col_count):
                            val = row_data[c_idx] if c_idx < len(row_data) else ""
                            cell = t.cell(r_idx, c_idx)
                            set_cell_margins(cell, top=120, bottom=120, left=150, right=150)
                            p = cell.paragraphs[0]
                            p.paragraph_format.space_before = Pt(2)
                            p.paragraph_format.space_after = Pt(2)
                            if r_idx == 0:
                                set_cell_background(cell, "0F2027")
                                run = p.add_run(re.sub(r'\*\*(.*?)\*\*', r'\1', val))
                                run.bold = True
                                run.font.name = 'Calibri'
                                run.font.size = Pt(10)
                                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                            else:
                                if r_idx % 2 == 1:
                                    set_cell_background(cell, "F9FAFC")
                                else:
                                    set_cell_background(cell, "FFFFFF")
                                run = p.add_run(re.sub(r'\*\*(.*?)\*\*', r'\1', val))
                                run.font.name = 'Calibri'
                                run.font.size = Pt(9.5)
                    doc.add_paragraph()
                table_rows = []
                i += 1
                continue

        # Headings
        if line.startswith('# '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(18)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run(line[2:].strip())
            run.bold = True
            run.font.name = 'Segoe UI'
            run.font.size = Pt(22)
            run.font.color.rgb = RGBColor(0x0F, 0x20, 0x27)
            i += 1
            continue

        if line.startswith('## '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(line[3:].strip())
            run.bold = True
            run.font.name = 'Segoe UI'
            run.font.size = Pt(15)
            run.font.color.rgb = RGBColor(0x20, 0x3A, 0x43)
            i += 1
            continue

        if line.startswith('### '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(3)
            run = p.add_run(line[4:].strip())
            run.bold = True
            run.font.name = 'Segoe UI'
            run.font.size = Pt(12.5)
            run.font.color.rgb = RGBColor(0x2C, 0x53, 0x64)
            i += 1
            continue

        if line.startswith('#### '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(line[5:].strip())
            run.bold = True
            run.font.name = 'Segoe UI'
            run.font.size = Pt(11)
            run.font.color.rgb = RGBColor(0x44, 0x44, 0x44)
            i += 1
            continue

        # Horizontal Rule
        if line.strip() in ['---', '***', '___']:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run('━' * 55)
            run.font.color.rgb = RGBColor(0xCC, 0xCC, 0xCC)
            run.font.size = Pt(8)
            i += 1
            continue

        # Bullet lists
        if line.strip().startswith('- ') or line.strip().startswith('* '):
            text = line.strip()[2:].strip()
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            # Parse bold parts
            parts = re.split(r'(\*\*.*?\*\*)', text)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    run = p.add_run(part[2:-2])
                    run.bold = True
                else:
                    p.add_run(part)
            i += 1
            continue

        # Numbered lists
        num_match = re.match(r'^\s*(\d+)\.\s+(.*)$', line)
        if num_match:
            text = num_match.group(2)
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            parts = re.split(r'(\*\*.*?\*\*)', text)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    run = p.add_run(part[2:-2])
                    run.bold = True
                else:
                    p.add_run(part)
            i += 1
            continue

        # Empty lines
        if not line.strip():
            i += 1
            continue

        # Standard paragraph
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        parts = re.split(r'(\*\*.*?\*\*)', line)
        for part in parts:
            if part.startswith('**') and part.endswith('**'):
                run = p.add_run(part[2:-2])
                run.bold = True
            else:
                p.add_run(part)
        i += 1

    doc.save(docx_path)
    print(f"Successfully converted {md_path} to {docx_path}")

if __name__ == '__main__':
    md_file = r"c:\Users\Ravitej Manu\Downloads\week1\GAME_FEATURE_DESIGN_DOCUMENT.md"
    docx_file = r"c:\Users\Ravitej Manu\Downloads\week1\Game_Feature_Design_Document_Week1.docx"
    create_docx_from_markdown(md_file, docx_file)
